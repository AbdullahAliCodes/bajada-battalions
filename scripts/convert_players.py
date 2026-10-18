#!/usr/bin/env python3
"""Convert the Team Padel Auction Excel export into data/players.json.

Expected workbook format (Google Form / Excel export):
  - First sheet, row 1 is headers
  - "Full name (Playtomic name)" column
  - "Playtomic Level" column

Output:
  [
    {"name": "Liam Day", "level": 4.5},
    ...
  ]

Usage:
  python3 -m venv .venv
  .venv/bin/pip install -r scripts/requirements.txt
  .venv/bin/python scripts/convert_players.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "Team Padel Auction WIP.xlsx"
OUTPUT = ROOT / "data" / "players.json"

NAME_HEADER = "Full name (Playtomic name)"
LEVEL_HEADER = "Playtomic Level"


def parse_level(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def convert(xlsx_path: Path = DEFAULT_XLSX, output_path: Path = OUTPUT) -> list[dict]:
    if not xlsx_path.exists():
        raise SystemExit(f"Excel file not found: {xlsx_path}")

    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    try:
        name_idx = headers.index(NAME_HEADER)
        level_idx = headers.index(LEVEL_HEADER)
    except ValueError as exc:
        raise SystemExit(
            f"Expected columns {NAME_HEADER!r} and {LEVEL_HEADER!r}. Found: {headers}"
        ) from exc

    players: list[dict] = []
    seen: set[str] = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        raw_name = row[name_idx]
        if not raw_name or not str(raw_name).strip():
            continue

        name = " ".join(str(raw_name).split())
        key = name.casefold()
        if key in seen:
            continue

        level = parse_level(row[level_idx])
        if level is None:
            continue

        seen.add(key)
        players.append({"name": name, "level": round(level, 2)})

    players.sort(key=lambda player: (-player["level"], player["name"].casefold()))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(players, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(players)} players to {output_path}")
    return players


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    convert(source)
